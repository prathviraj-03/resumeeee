from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import asyncio
import uuid
import asyncpg
from urllib.parse import urlparse
from dotenv import load_dotenv

def add_skills(doc):
    doc.add_heading('Skills', level=1)
    doc.add_paragraph('{% for skill in skills %}{{ skill }}{% if not loop.last %}, {% endif %}{% endfor %}')

def add_experience(doc):
    doc.add_heading('Experience', level=1)
    doc.add_paragraph('{% for exp in experience %}')
    
    p = doc.add_paragraph('{{ exp.title }} at {{ exp.company }}')
    p.runs[0].bold = True
    
    doc.add_paragraph('{{ exp.duration }}')
    doc.add_paragraph('{{ exp.description }}')
    doc.add_paragraph('{% endfor %}')

def add_education(doc):
    doc.add_heading('Education', level=1)
    doc.add_paragraph('{% for edu in education %}')
    
    p = doc.add_paragraph('{{ edu.degree }} - {{ edu.institution }}')
    p.runs[0].bold = True
    
    doc.add_paragraph('{{ edu.year }}')
    doc.add_paragraph('{% endfor %}')

def create_fresher_modern():
    doc = Document()
    name = doc.add_paragraph('{{ full_name }}')
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.runs[0]
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 102, 204)

    contact = doc.add_paragraph('{{ phone_number }} | {{ linkedin_url }} | {{ portfolio_url }}')
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(20)

    doc.add_heading('Professional Summary', level=1)
    doc.add_paragraph('{{ summary }}')

    add_education(doc)
    add_skills(doc)
    add_experience(doc)

    os.makedirs('seeds/templates', exist_ok=True)
    filepath = 'seeds/templates/Fresher_Modern.docx'
    doc.save(filepath)
    return filepath, "Fresher Modern", "Clean, blue-themed modern layout for entry-level"

def create_fresher_classic():
    doc = Document()
    name = doc.add_paragraph('{{ full_name }}')
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = name.runs[0]
    run.bold = True
    run.font.size = Pt(22)

    contact = doc.add_paragraph('{{ phone_number }} | {{ linkedin_url }} | {{ portfolio_url }}')
    contact.alignment = WD_ALIGN_PARAGRAPH.LEFT
    contact.paragraph_format.space_after = Pt(15)

    doc.add_heading('Objective', level=1)
    doc.add_paragraph('{{ summary }}')

    add_education(doc)
    add_skills(doc)
    add_experience(doc)

    os.makedirs('seeds/templates', exist_ok=True)
    filepath = 'seeds/templates/Fresher_Classic.docx'
    doc.save(filepath)
    return filepath, "Fresher Classic", "Traditional black-and-white academic layout"

def create_experienced_professional():
    doc = Document()
    name = doc.add_paragraph('{{ full_name }}')
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.runs[0]
    run.bold = True
    run.font.size = Pt(26)
    
    role = doc.add_paragraph('{{ target_role }}')
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role.runs[0].italic = True

    contact = doc.add_paragraph('{{ phone_number }} | {{ linkedin_url }} | {{ portfolio_url }}')
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(20)

    doc.add_heading('Professional Summary', level=1)
    doc.add_paragraph('{{ summary }}')

    add_experience(doc)
    add_skills(doc)
    add_education(doc)

    os.makedirs('seeds/templates', exist_ok=True)
    filepath = 'seeds/templates/Experienced_Professional.docx'
    doc.save(filepath)
    return filepath, "Experienced Professional", "Focus on experience and achievements"

def create_experienced_executive():
    doc = Document()
    name = doc.add_paragraph('{{ full_name }}')
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = name.runs[0]
    run.bold = True
    run.font.size = Pt(28)
    
    role = doc.add_paragraph('{{ target_role }}')
    role.alignment = WD_ALIGN_PARAGRAPH.LEFT
    role.runs[0].font.size = Pt(16)
    role.runs[0].italic = True

    contact = doc.add_paragraph('Phone: {{ phone_number }} | LinkedIn: {{ linkedin_url }}')
    contact.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_heading('Executive Profile', level=1)
    doc.add_paragraph('{{ summary }}')

    add_skills(doc)
    add_experience(doc)
    add_education(doc)

    os.makedirs('seeds/templates', exist_ok=True)
    filepath = 'seeds/templates/Experienced_Executive.docx'
    doc.save(filepath)
    return filepath, "Experienced Executive", "Bold and spacious layout for senior roles"

async def seed_database(templates_data):
    import jwt
    import datetime
    import httpx
    
    # 1. Create a dummy JWT token for a system user or admin
    secret = "change-me-in-production"
    payload = {
        "userId": "system-seeder-001",
        "email": "admin@resumeforge.com",
        "role": "admin",
        "exp": datetime.datetime.utcnow() + datetime.timedelta(hours=1)
    }
    token = jwt.encode(payload, secret, algorithm="HS256")
    headers = {"Authorization": f"Bearer {token}"}
    
    # The profile service is exposed on port 8003 locally via docker
    profile_api_url = "http://localhost:8003/api/templates/"
    
    print(f"Uploading to {profile_api_url} ...")
    
    async with httpx.AsyncClient() as client:
        for filepath, name, desc in templates_data:
            abs_path = os.path.abspath(filepath)
            
            with open(abs_path, "rb") as f:
                files = {"file": (os.path.basename(abs_path), f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
                data = {
                    "name": name,
                    "description": desc,
                    "is_default": "true" if "Modern" in name else "false"
                }
                
                try:
                    response = await client.post(profile_api_url, headers=headers, data=data, files=files)
                    if response.status_code in (200, 201):
                        print(f"[OK] Successfully seeded: {name}")
                    else:
                        print(f"[FAIL] Failed to seed {name}: {response.status_code}")
                except Exception as e:
                    print(f"[ERROR] Error uploading {name}: {type(e).__name__}")

if __name__ == "__main__":
    templates = [
        create_fresher_modern(),
        create_fresher_classic(),
        create_experienced_professional(),
        create_experienced_executive()
    ]
    
    # Run async seeder
    asyncio.run(seed_database(templates))
